"""Excel workbook generation and formatting."""

from __future__ import annotations

from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from reconforge.utils.time import utc_now_text

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
LOW_FILL = PatternFill("solid", fgColor="E2F0D9")
MEDIUM_FILL = PatternFill("solid", fgColor="FFF2CC")
HIGH_FILL = PatternFill("solid", fgColor="FCE4D6")
CRITICAL_FILL = PatternFill("solid", fgColor="F4CCCC")
TOTAL_FILL = PatternFill("solid", fgColor="D9EAF7")


def autosize_columns(workbook_path: Path) -> None:
    """Apply common workbook styling."""

    workbook = load_workbook(workbook_path)
    for worksheet in workbook.worksheets:
        if worksheet.max_row == 0:
            continue
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for cell in worksheet[1]:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center")

        for column_cells in worksheet.columns:
            values = [str(cell.value) for cell in column_cells if cell.value is not None]
            width = min(max([len(value) for value in values], default=10) + 2, 48)
            worksheet.column_dimensions[get_column_letter(column_cells[0].column)].width = width

        risk_column = None
        for index, cell in enumerate(worksheet[1], start=1):
            if str(cell.value).lower() == "risk_score":
                risk_column = index
                break
        if risk_column is not None and worksheet.max_row > 1:
            col = get_column_letter(risk_column)
            rng = f"{col}2:{col}{worksheet.max_row}"
            worksheet.conditional_formatting.add(rng, CellIsRule(operator="lessThanOrEqual", formula=["30"], fill=LOW_FILL))
            worksheet.conditional_formatting.add(rng, CellIsRule(operator="between", formula=["31", "60"], fill=MEDIUM_FILL))
            worksheet.conditional_formatting.add(rng, CellIsRule(operator="between", formula=["61", "80"], fill=HIGH_FILL))
            worksheet.conditional_formatting.add(rng, CellIsRule(operator="greaterThan", formula=["80"], fill=CRITICAL_FILL))

    workbook.save(workbook_path)


def write_excel_workbook(
    sheets: dict[str, pd.DataFrame],
    output_path: Path | str,
    metadata: dict[str, Any] | None = None,
) -> Path:
    """Write report sheets to a formatted Excel workbook."""

    workbook_path = Path(output_path)
    workbook_path.parent.mkdir(parents=True, exist_ok=True)

    normalized_sheets = {name[:31]: frame.copy() for name, frame in sheets.items()}
    if metadata:
        metadata_frame = pd.DataFrame([metadata])
        normalized_sheets = {"Report Parameters": metadata_frame, **normalized_sheets}

    with pd.ExcelWriter(workbook_path, engine="openpyxl") as writer:
        for name, frame in normalized_sheets.items():
            safe_frame = frame.copy()
            for column in safe_frame.columns:
                if pd.api.types.is_datetime64_any_dtype(safe_frame[column]):
                    safe_frame[column] = safe_frame[column].dt.strftime("%Y-%m-%d")
            safe_frame.to_excel(writer, sheet_name=name, index=False)

    autosize_columns(workbook_path)
    return workbook_path


def add_summary_chart(workbook_path: Path | str, sheet_name: str, title: str) -> None:
    """Add a simple bar chart to a two-column summary sheet."""

    path = Path(workbook_path)
    workbook = load_workbook(path)
    if sheet_name not in workbook.sheetnames:
        workbook.save(path)
        return
    worksheet = workbook[sheet_name]
    if worksheet.max_row < 2 or worksheet.max_column < 2:
        workbook.save(path)
        return

    chart = BarChart()
    chart.title = title
    chart.y_axis.title = "Count"
    chart.x_axis.title = "Category"
    data = Reference(worksheet, min_col=2, min_row=1, max_row=worksheet.max_row)
    categories = Reference(worksheet, min_col=1, min_row=2, max_row=worksheet.max_row)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    chart.height = 7
    chart.width = 12
    worksheet.add_chart(chart, "D2")
    workbook.save(path)


def audit_metadata(company_name: str, report_title: str, currency: str) -> dict[str, str]:
    """Create standard report metadata."""

    return {
        "company_name": company_name,
        "report_title": report_title,
        "output_currency": currency,
        "generated_at": utc_now_text(),
        "tool": "ReconForge ERP",
    }
