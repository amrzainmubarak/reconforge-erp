from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from typer.testing import CliRunner

from reconforge.cli import app
from reconforge.periods import _amount_value
from reconforge.review.state import load_review_state, save_review_state, update_review_status

runner = CliRunner()


def _write_period(path: Path, rows: list[dict[str, object]]) -> None:
    path.mkdir(parents=True, exist_ok=True)
    pd.DataFrame(rows).to_csv(path / "stock_gl_all_exceptions.csv", index=False)


def test_compare_periods_detects_new_recurring_and_resolved(tmp_path: Path) -> None:
    jan = tmp_path / "jan"
    feb = tmp_path / "feb"
    _write_period(
        jan,
        [
            {"exception_id": "EXC-001", "exception_type": "stock_without_gl", "risk_level": "High", "source_document": "SO-1", "total_cost": 100},
            {"exception_id": "EXC-002", "exception_type": "gl_without_stock", "risk_level": "Medium", "reference": "JE-1", "amount": 50},
        ],
    )
    _write_period(
        feb,
        [
            {"exception_id": "EXC-001", "exception_type": "stock_without_gl", "risk_level": "High", "source_document": "SO-1", "total_cost": 100},
            {"exception_id": "EXC-003", "exception_type": "stock_without_gl", "risk_level": "Critical", "source_document": "SO-2", "total_cost": 250},
        ],
    )
    output = tmp_path / "compare"
    result = runner.invoke(app, ["compare", "periods", "--inputs", str(jan), str(feb), "--output", str(output)])
    assert result.exit_code == 0
    payload = json.loads((output / "period_comparison.json").read_text(encoding="utf-8"))
    summary = {row["metric"]: row["value"] for row in payload["summary"]}
    assert summary["new_exceptions"] == 1
    assert summary["recurring_exceptions"] == 1
    assert summary["resolved_exceptions"] == 1
    assert payload["trend"]["chart_data"]["new"] == [2, 1]
    assert payload["trend"]["chart_data"]["recurring"] == [0, 1]
    assert "top_recurring_themes" in payload["trend"]
    assert (output / "period_comparison.xlsx").exists()
    assert (output / "period_comparison.html").exists()
    assert (output / "period_comparison.md").exists()
    html = (output / "period_comparison.html").read_text(encoding="utf-8")
    assert "Trend Summary" in html
    workbook = load_workbook(output / "period_comparison.xlsx", read_only=True)
    assert "Trend Summary" in workbook.sheetnames
    assert "Top Recurring Themes" in workbook.sheetnames


def test_period_fingerprint_amount_uses_decimal_rounding() -> None:
    row = pd.Series({"amount": "100.005"})

    assert _amount_value(row) == "100.01"


def test_compare_periods_matches_reordered_synthetic_exception_ids_by_fingerprint(tmp_path: Path) -> None:
    jan = tmp_path / "jan"
    feb = tmp_path / "feb"
    _write_period(
        jan,
        [
            {"exception_type": "stock_without_gl", "source_document": "SO-1", "product_code": "P-1", "date": "2026-01-31", "total_cost": 100},
            {"exception_type": "gl_without_stock", "reference": "JE-2", "account_code": "5000", "date": "2026-01-31", "amount": 250},
        ],
    )
    _write_period(
        feb,
        [
            {"exception_type": "gl_without_stock", "reference": "JE-2", "account_code": "5000", "date": "2026-01-31", "amount": 250},
            {"exception_type": "stock_without_gl", "source_document": "SO-1", "product_code": "P-1", "date": "2026-01-31", "total_cost": 100},
        ],
    )
    output = tmp_path / "compare"
    result = runner.invoke(app, ["compare", "periods", "--inputs", str(jan), str(feb), "--output", str(output)])
    assert result.exit_code == 0
    payload = json.loads((output / "period_comparison.json").read_text(encoding="utf-8"))
    summary = {row["metric"]: row["value"] for row in payload["summary"]}
    assert summary["recurring_exceptions"] == 2
    assert summary["new_exceptions"] == 0
    assert summary["resolved_exceptions"] == 0


def test_compare_periods_does_not_match_same_row_position_exception_id_by_itself(tmp_path: Path) -> None:
    jan = tmp_path / "jan"
    feb = tmp_path / "feb"
    _write_period(
        jan,
        [{"exception_id": "EXC-0001", "exception_type": "stock_without_gl", "source_document": "SO-OLD", "product_code": "P-1", "total_cost": 100}],
    )
    _write_period(
        feb,
        [{"exception_id": "EXC-0001", "exception_type": "stock_without_gl", "source_document": "SO-NEW", "product_code": "P-2", "total_cost": 200}],
    )
    output = tmp_path / "compare"
    result = runner.invoke(app, ["compare", "periods", "--inputs", str(jan), str(feb), "--output", str(output)])
    assert result.exit_code == 0
    payload = json.loads((output / "period_comparison.json").read_text(encoding="utf-8"))
    summary = {row["metric"]: row["value"] for row in payload["summary"]}
    assert summary["recurring_exceptions"] == 0
    assert summary["new_exceptions"] == 1
    assert summary["resolved_exceptions"] == 1


def test_compare_periods_includes_review_state_categories(tmp_path: Path) -> None:
    jan = tmp_path / "jan"
    feb = tmp_path / "feb"
    _write_period(jan, [{"exception_id": "EXC-001", "exception_type": "stock_without_gl", "risk_level": "High"}])
    _write_period(
        feb,
        [
            {"exception_id": "EXC-001", "exception_type": "stock_without_gl", "risk_level": "High"},
            {"exception_id": "EXC-002", "exception_type": "gl_without_stock", "risk_level": "Critical"},
        ],
    )
    state = load_review_state(feb / "review_state.json")
    update_review_status("EXC-001", "Escalated", state, escalation_owner="Controller")
    update_review_status("EXC-002", "Accepted Risk", state, accepted_risk_reason="Documented timing")
    save_review_state(feb / "review_state.json", state)
    output = tmp_path / "compare"
    result = runner.invoke(app, ["compare", "periods", "--inputs", str(jan), "--inputs", str(feb), "--output", str(output)])
    assert result.exit_code == 0
    payload = json.loads((output / "period_comparison.json").read_text(encoding="utf-8"))
    assert len(payload["escalated_exceptions"]) == 1
    assert len(payload["accepted_risk_items"]) == 1


def test_compare_periods_missing_folder_handled_cleanly(tmp_path: Path) -> None:
    jan = tmp_path / "jan"
    _write_period(jan, [{"exception_id": "EXC-001", "exception_type": "stock_without_gl"}])
    result = runner.invoke(app, ["compare", "periods", "--inputs", str(jan), "--inputs", str(tmp_path / "missing"), "--output", str(tmp_path / "out")])
    assert result.exit_code == 1
    assert "Missing period folder" in result.output
