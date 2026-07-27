from __future__ import annotations

import json
from decimal import Decimal
from pathlib import Path

import pytest
from jsonschema import Draft202012Validator
from openpyxl import load_workbook
from typer.testing import CliRunner

from reconforge.cli import app
from reconforge.periods import (
    compare_period_outputs,
    read_period_comparison,
    verify_period_comparison_payload,
)
from reconforge.review.state import collect_exception_frame, export_review_register
from reconforge.utils.money import (
    LEGACY_FINANCIAL_INPUT_POLICY,
    STRICT_FINANCIAL_INPUT_POLICY,
    LegacyFinancialInputWarning,
)

runner = CliRunner()


def _write_period(path: Path, *, amount: str, reference: str = "DOC-1") -> None:
    path.mkdir(parents=True, exist_ok=True)
    (path / "stock_gl_all_exceptions.csv").write_text(
        "exception_type,source_document,total_cost\n"
        f"stock_without_gl,{reference},{amount}\n",
        encoding="utf-8",
    )


def _summary(payload: dict[str, object]) -> dict[str, object]:
    rows = payload["summary"]
    assert isinstance(rows, list)
    return {
        str(row["metric"]): row["value"]
        for row in rows
        if isinstance(row, dict)
    }


def _legacy_payload(current: dict[str, object]) -> dict[str, object]:
    keys = {
        "periods",
        "summary",
        "period_counts",
        "trend",
        "new_exceptions",
        "recurring_exceptions",
        "resolved_exceptions",
        "escalated_exceptions",
        "accepted_risk_items",
    }
    return {key: current[key] for key in keys}


def test_strict_period_reader_preserves_rounding_boundary_legacy_collapses(
    tmp_path: Path,
) -> None:
    first = tmp_path / "first"
    second = tmp_path / "second"
    _write_period(first, amount="9007199254740992.004")
    _write_period(second, amount="9007199254740992.005")

    strict = compare_period_outputs(
        [first, second],
        tmp_path / "strict",
        financial_input_policy=STRICT_FINANCIAL_INPUT_POLICY,
    )
    with pytest.warns(LegacyFinancialInputWarning, match="deprecated"):
        legacy = compare_period_outputs(
            [first, second],
            tmp_path / "legacy",
            financial_input_policy=LEGACY_FINANCIAL_INPUT_POLICY,
        )
    strict_payload = json.loads(strict.json_path.read_text(encoding="utf-8"))
    legacy_payload = json.loads(legacy.json_path.read_text(encoding="utf-8"))

    assert _summary(strict_payload)["recurring_exceptions"] == 0
    assert _summary(strict_payload)["new_exceptions"] == 1
    assert _summary(strict_payload)["resolved_exceptions"] == 1
    assert _summary(legacy_payload)["recurring_exceptions"] == 1
    assert strict_payload["decision_digest"] != legacy_payload["decision_digest"]


def test_strict_review_ingress_preserves_exact_text_and_invalid_absence(
    tmp_path: Path,
) -> None:
    source = tmp_path / "source"
    source.mkdir()
    (source / "stock_gl_all_exceptions.csv").write_text(
        "exception_type,source_document,total_cost\n"
        "stock_without_gl,DOC-1,9007199254740992.004\n",
        encoding="utf-8",
    )
    (source / "workorders_all_exceptions.csv").write_text(
        "exception_type,source_document,total_cost\n"
        "stock_without_gl,DOC-2,malformed\n",
        encoding="utf-8",
    )

    strict = collect_exception_frame(
        source,
        financial_input_policy=STRICT_FINANCIAL_INPUT_POLICY,
    )
    legacy = collect_exception_frame(
        source,
        financial_input_policy=LEGACY_FINANCIAL_INPUT_POLICY,
    )

    assert strict.loc[0, "amount_impact"] == Decimal("9007199254740992.004")
    assert legacy.loc[0, "amount_impact"] == Decimal("9007199254740992.0")
    assert strict.loc[1, "amount_impact"] is None


def test_review_register_records_policy_without_moving_compatibility_sheet(
    tmp_path: Path,
) -> None:
    source = tmp_path / "source"
    _write_period(source, amount="10")
    output = export_review_register(
        source,
        tmp_path / "review.xlsx",
        financial_input_policy=STRICT_FINANCIAL_INPUT_POLICY,
    )

    workbook = load_workbook(output, read_only=True)
    assert workbook.sheetnames[:2] == ["Review Register", "Report Parameters"]
    parameters = workbook["Report Parameters"]
    assert parameters["A1"].value == "financial_input_policy"
    assert parameters["A2"].value == STRICT_FINANCIAL_INPUT_POLICY


def test_strict_period_reader_does_not_turn_invalid_amount_into_zero(
    tmp_path: Path,
) -> None:
    first = tmp_path / "first"
    second = tmp_path / "second"
    _write_period(first, amount="malformed")
    _write_period(second, amount="0")

    strict = compare_period_outputs(
        [first, second],
        tmp_path / "strict",
        financial_input_policy=STRICT_FINANCIAL_INPUT_POLICY,
    )
    legacy = compare_period_outputs(
        [first, second],
        tmp_path / "legacy",
        financial_input_policy=LEGACY_FINANCIAL_INPUT_POLICY,
    )
    strict_payload = json.loads(strict.json_path.read_text(encoding="utf-8"))
    legacy_payload = json.loads(legacy.json_path.read_text(encoding="utf-8"))

    assert _summary(strict_payload)["recurring_exceptions"] == 0
    assert "amount=invalid" in strict_payload["resolved_exceptions"][0][
        "comparison_key"
    ]
    assert _summary(legacy_payload)["recurring_exceptions"] == 1


def test_current_period_artifact_schema_digests_and_surfaces(tmp_path: Path) -> None:
    first = tmp_path / "first"
    second = tmp_path / "second"
    _write_period(first, amount="10")
    _write_period(second, amount="10")
    artifacts = compare_period_outputs(
        [first, second],
        tmp_path / "output",
        financial_input_policy=STRICT_FINANCIAL_INPUT_POLICY,
    )
    payload = json.loads(artifacts.json_path.read_text(encoding="utf-8"))
    schema = json.loads(
        Path("docs/schemas/period_comparison.schema.json").read_text(
            encoding="utf-8"
        )
    )

    Draft202012Validator.check_schema(schema)
    Draft202012Validator(schema).validate(payload)
    document = read_period_comparison(artifacts.json_path)
    verify_period_comparison_payload(payload, period_paths=[first, second])
    assert document.schema_version == 2
    assert document.verification_status == "verified"
    assert payload["financial_input_policy"] == STRICT_FINANCIAL_INPUT_POLICY
    assert payload["comparison_policy"]["invalid_amount_policy"] == (
        "explicit-invalid-or-missing-v2"
    )
    assert payload["input_periods"][0]["files"][0]["name"] == (
        "stock_gl_all_exceptions.csv"
    )
    workbook = load_workbook(artifacts.workbook_path, read_only=True)
    parameters = workbook["Report Parameters"]
    assert parameters["A1"].value == "financial_input_policy"
    assert parameters["A2"].value == STRICT_FINANCIAL_INPUT_POLICY
    assert STRICT_FINANCIAL_INPUT_POLICY in artifacts.html_path.read_text(
        encoding="utf-8"
    )
    assert STRICT_FINANCIAL_INPUT_POLICY in artifacts.markdown_path.read_text(
        encoding="utf-8"
    )


def test_period_reader_preserves_historical_unversioned_v1(tmp_path: Path) -> None:
    first = tmp_path / "first"
    second = tmp_path / "second"
    _write_period(first, amount="10")
    _write_period(second, amount="10")
    artifacts = compare_period_outputs(
        [first, second],
        tmp_path / "output",
        financial_input_policy=STRICT_FINANCIAL_INPUT_POLICY,
    )
    current = json.loads(artifacts.json_path.read_text(encoding="utf-8"))
    legacy = _legacy_payload(current)
    legacy_path = tmp_path / "legacy.json"
    legacy_path.write_text(json.dumps(legacy), encoding="utf-8")
    schema = json.loads(
        Path("docs/schemas/period_comparison.schema.json").read_text(
            encoding="utf-8"
        )
    )

    Draft202012Validator(schema).validate(legacy)
    document = read_period_comparison(legacy_path)
    assert document.schema_version == 1
    assert document.verification_status == "legacy-unverified"


def test_period_verifier_detects_decision_artifact_and_source_tampering(
    tmp_path: Path,
) -> None:
    first = tmp_path / "first"
    second = tmp_path / "second"
    _write_period(first, amount="10")
    _write_period(second, amount="10")
    artifacts = compare_period_outputs(
        [first, second],
        tmp_path / "output",
        financial_input_policy=STRICT_FINANCIAL_INPUT_POLICY,
    )
    payload = json.loads(artifacts.json_path.read_text(encoding="utf-8"))
    changed_decision = json.loads(json.dumps(payload))
    changed_decision["summary"][0]["value"] += 1
    with pytest.raises(ValueError, match="decision digest"):
        verify_period_comparison_payload(changed_decision)

    changed_path = json.loads(json.dumps(payload))
    changed_path["periods"][0] = "relocated"
    with pytest.raises(ValueError, match="artifact digest"):
        verify_period_comparison_payload(changed_path)

    _write_period(first, amount="11")
    with pytest.raises(ValueError, match="input fingerprint"):
        verify_period_comparison_payload(payload, period_paths=[first, second])


def test_period_decision_digest_is_path_independent_for_identical_inputs(
    tmp_path: Path,
) -> None:
    first_a = tmp_path / "a" / "first"
    second_a = tmp_path / "a" / "second"
    first_b = tmp_path / "b" / "first"
    second_b = tmp_path / "b" / "second"
    for path in (first_a, second_a, first_b, second_b):
        _write_period(path, amount="10")

    output_a = compare_period_outputs(
        [first_a, second_a],
        tmp_path / "output-a",
        financial_input_policy=STRICT_FINANCIAL_INPUT_POLICY,
    )
    output_b = compare_period_outputs(
        [first_b, second_b],
        tmp_path / "output-b",
        financial_input_policy=STRICT_FINANCIAL_INPUT_POLICY,
    )
    payload_a = json.loads(output_a.json_path.read_text(encoding="utf-8"))
    payload_b = json.loads(output_b.json_path.read_text(encoding="utf-8"))

    assert payload_a["decision_digest"] == payload_b["decision_digest"]
    assert payload_a["artifact_digest"] != payload_b["artifact_digest"]


def test_cli_period_comparison_selects_strict_policy(tmp_path: Path) -> None:
    first = tmp_path / "first"
    second = tmp_path / "second"
    output = tmp_path / "output"
    _write_period(first, amount="0.0049999999999999999")
    _write_period(second, amount="0.005")

    result = runner.invoke(
        app,
        [
            "compare",
            "periods",
            "--inputs",
            str(first),
            "--inputs",
            str(second),
            "--output",
            str(output),
        ],
    )

    assert result.exit_code == 0
    payload = json.loads((output / "period_comparison.json").read_text(encoding="utf-8"))
    assert payload["financial_input_policy"] == STRICT_FINANCIAL_INPUT_POLICY
    assert _summary(payload)["recurring_exceptions"] == 0


def test_unknown_period_policy_fails_before_path_or_output(tmp_path: Path) -> None:
    missing = tmp_path / "987654321.01"
    output = tmp_path / "output"

    with pytest.raises(ValueError, match="unsupported financial input policy") as captured:
        compare_period_outputs(
            [missing, missing],
            output,
            financial_input_policy="unknown-v9",  # type: ignore[arg-type]
        )

    assert "987654321.01" not in str(captured.value)
    assert not output.exists()
