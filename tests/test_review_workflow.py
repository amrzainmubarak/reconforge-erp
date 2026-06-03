from __future__ import annotations

import json
import re
from collections.abc import Callable
from pathlib import Path
from typing import Any, cast

import pandas as pd
import pytest
from fastapi.routing import APIRoute
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from typer.testing import CliRunner

from reconforge.cli import app
from reconforge.evidence.binder import generate_evidence_binder
from reconforge.review.state import (
    export_review_register,
    get_review_status,
    load_review_state,
    merge_review_state_with_exceptions,
    save_review_state,
    update_review_status,
)
from reconforge.studio.app import _table, create_studio_app

runner = CliRunner()


def _write_exception_output(output_dir: Path) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    pd.DataFrame(
        [
            {
                "exception_type": "stock_without_gl",
                "risk_level": "Critical",
                "risk_score": 90,
                "work_order": "WO-1",
                "product_code": "P-1",
                "total_cost": 1250,
            },
            {
                "exception_type": "gl_without_stock",
                "risk_level": "High",
                "risk_score": 70,
                "work_order": "WO-2",
                "product_code": "P-2",
                "total_cost": 300,
            },
        ],
    ).to_csv(output_dir / "stock_gl_all_exceptions.csv", index=False)


def _render_studio_exceptions(output_dir: Path, **kwargs: Any) -> str:
    studio = create_studio_app("examples/sample_data", output_dir)
    for route in studio.routes:
        if isinstance(route, APIRoute) and route.path == "/exceptions":
            endpoint = cast(Callable[..., str], route.endpoint)
            return endpoint(**kwargs)
    raise AssertionError("Studio exceptions route not found")


def _csrf_token(html: str) -> str:
    match = re.search(r'name="csrf_token" value="([^"]+)"', html)
    assert match is not None
    return match.group(1)


def test_review_state_load_save_and_default_status(tmp_path: Path) -> None:
    path = tmp_path / "review_state.json"
    state = load_review_state(path)
    assert state == {}
    update_review_status("EXC-0001", "Under Review", state, reviewer="Amr", note="Checking source records")
    save_review_state(path, state)
    loaded = load_review_state(path)
    assert get_review_status("EXC-0001", loaded) == "Under Review"
    assert get_review_status("EXC-404", loaded) == "New"


def test_review_state_rejects_invalid_status() -> None:
    with pytest.raises(ValueError):
        update_review_status("EXC-0001", "Done", {})


def test_malformed_review_state_json_falls_back_safely(tmp_path: Path) -> None:
    path = tmp_path / "review_state.json"
    path.write_text("{bad json", encoding="utf-8")
    assert load_review_state(path) == {}


def test_review_state_merge_with_exceptions_preserves_unknown_entries() -> None:
    state = {
        "EXC-0001": {"exception_id": "EXC-0001", "status": "Resolved", "reviewer": "Amr"},
        "EXC-9999": {"exception_id": "EXC-9999", "status": "Escalated", "reviewer": "Finance"},
    }
    frame = pd.DataFrame([{"exception_type": "stock_without_gl", "risk_score": 90, "total_cost": 100}])
    merged = merge_review_state_with_exceptions(frame, state)
    assert merged.loc[0, "exception_id"] == "EXC-0001"
    assert merged.loc[0, "status"] == "Resolved"
    assert "EXC-9999" in state


def test_cli_review_set_status_list_and_export(tmp_path: Path) -> None:
    output = tmp_path / "output"
    _write_exception_output(output)
    set_result = runner.invoke(
        app,
        [
            "review",
            "set-status",
            "--input",
            str(output),
            "--exception-id",
            "EXC-0001",
            "--status",
            "Under Review",
            "--reviewer",
            "Amr",
            "--note",
            "Checking source records",
        ],
    )
    assert set_result.exit_code == 0
    assert (output / "review_state.json").exists()

    list_result = runner.invoke(app, ["review", "list", "--input", str(output)])
    assert list_result.exit_code == 0
    assert "Under Review" in list_result.output

    register_path = output / "review_register.xlsx"
    export_result = runner.invoke(app, ["review", "export", "--input", str(output), "--output", str(register_path)])
    assert export_result.exit_code == 0
    assert register_path.exists()


def test_review_register_export_contains_review_fields(tmp_path: Path) -> None:
    output = tmp_path / "output"
    _write_exception_output(output)
    state = load_review_state(output / "review_state.json")
    update_review_status("EXC-0001", "Escalated", state, reviewer="Controller", escalation_owner="CFO")
    save_review_state(output / "review_state.json", state)
    path = export_review_register(output, output / "review_register.xlsx")
    frame = pd.read_excel(path)
    assert "status" in frame.columns
    assert frame.loc[0, "status"] == "Escalated"


def test_evidence_register_includes_review_status_when_state_exists(tmp_path: Path) -> None:
    output = tmp_path / "output"
    _write_exception_output(output)
    state = load_review_state(output / "review_state.json")
    update_review_status("EXC-0001", "Accepted Risk", state, reviewer="Audit", accepted_risk_reason="Documented timing")
    save_review_state(output / "review_state.json", state)

    artifacts = generate_evidence_binder(output, tmp_path / "evidence")
    assert artifacts
    summary = (artifacts[0].folder / "summary.md").read_text(encoding="utf-8")
    assert "Accepted Risk" in summary
    workbook = load_workbook(tmp_path / "evidence" / "evidence_register.xlsx")
    headers = [cell.value for cell in next(workbook["Evidence Register"].iter_rows(max_row=1))]
    assert "status" in headers
    status_index = headers.index("status") + 1
    assert workbook["Evidence Register"].cell(row=2, column=status_index).value == "Accepted Risk"
    index_payload = json.loads((tmp_path / "evidence" / "evidence_index.json").read_text(encoding="utf-8"))
    assert index_payload["cases"][0]["review_status"] == "Accepted Risk"


def test_studio_filters_render_and_status_filter_works(tmp_path: Path) -> None:
    state = load_review_state(tmp_path / "review_state.json")
    update_review_status("EXC-0001", "Escalated", state, reviewer="Amr")
    save_review_state(tmp_path / "review_state.json", state)
    html = _render_studio_exceptions(tmp_path, status="Escalated")
    assert "Review status" in html
    assert "Escalated" in html


def test_studio_valid_status_update_persists_to_review_state(tmp_path: Path) -> None:
    client = TestClient(create_studio_app("examples/sample_data", tmp_path))
    token = _csrf_token(client.get("/exceptions").text)
    response = client.post(
        "/exceptions/update-review",
        data={
            "csrf_token": token,
            "exception_id": "EXC-0001",
            "status": "Under Review",
            "reviewer": "Controller",
            "note": "Checking source evidence",
        },
    )
    assert response.status_code == 200
    assert "Review updated for EXC-0001 as Under Review." in response.text
    state = load_review_state(tmp_path / "review_state.json")
    assert state["EXC-0001"]["status"] == "Under Review"
    assert state["EXC-0001"]["reviewer"] == "Controller"


def test_studio_invalid_status_rejected(tmp_path: Path) -> None:
    client = TestClient(create_studio_app("examples/sample_data", tmp_path))
    token = _csrf_token(client.get("/exceptions").text)
    response = client.post(
        "/exceptions/update-review",
        data={"csrf_token": token, "exception_id": "EXC-0001", "status": "Done"},
    )
    assert response.status_code == 200
    assert "Invalid review status" in response.text
    assert not (tmp_path / "review_state.json").exists()


def test_studio_review_update_escapes_reviewer_and_note(tmp_path: Path) -> None:
    client = TestClient(create_studio_app("examples/sample_data", tmp_path))
    token = _csrf_token(client.get("/exceptions").text)
    response = client.post(
        "/exceptions/update-review",
        data={
            "csrf_token": token,
            "exception_id": "EXC-0001",
            "status": "Under Review",
            "reviewer": "<script>alert(1)</script>",
            "note": "<b>needs review</b>",
        },
    )
    assert "<script>alert(1)</script>" not in response.text
    assert "<b>needs review</b>" not in response.text
    assert "&lt;script&gt;alert(1)&lt;/script&gt;" in response.text
    assert "&lt;b&gt;needs review&lt;/b&gt;" in response.text


def test_studio_filters_reflect_updated_status(tmp_path: Path) -> None:
    client = TestClient(create_studio_app("examples/sample_data", tmp_path))
    token = _csrf_token(client.get("/exceptions").text)
    client.post(
        "/exceptions/update-review",
        data={"csrf_token": token, "exception_id": "EXC-0001", "status": "Accepted Risk", "accepted_risk_reason": "Timing difference documented"},
    )
    filtered = client.get("/exceptions", params={"status": "Accepted Risk"}).text
    assert "Accepted Risk" in filtered
    assert "Timing difference documented" in filtered
    assert "No exceptions match the current filters." not in filtered


def test_studio_empty_filter_result_message(tmp_path: Path) -> None:
    html = _render_studio_exceptions(tmp_path, search="no-such-exception-text")
    assert "No exceptions match the current filters." in html


def test_studio_table_escapes_dataframe_values() -> None:
    html = _table(pd.DataFrame([{"exception_type": "<script>alert(1)</script>", "risk_score": 90}]))
    assert "<script>alert(1)</script>" not in html
    assert "&lt;script&gt;alert(1)&lt;/script&gt;" in html
